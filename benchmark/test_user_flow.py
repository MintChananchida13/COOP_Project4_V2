import unittest
import importlib.util
import io
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import Mock, patch

import user_flow


class UserFlowTests(unittest.TestCase):
    def test_canonical_null_zero_and_backend_source_priority(self):
        timing = {"image_siglip_runtime_ms": None, "siglip_inference_ms": 99,
                  "total_detection_ms": 10, "backend_runtime_ms": 0,
                  "benchmark": {"image_siglip_runtime_ms": 88}}
        values = user_flow.timing_values([[{"metric": "backend_runtime_ms", "valueMs": 50},
                                          {"metric": "total_runtime_ms", "valueMs": None}]], {"debug": {"timing": timing}})
        self.assertIsNone(values["image_siglip_runtime_ms"])
        self.assertEqual(values["backend_runtime_ms"], 0)
        self.assertEqual(values["total_runtime_ms"], 10)
        for invalid in (True, float("inf"), float("nan"), -1, "20"):
            self.assertIsNone(user_flow.valid_duration(invalid))

    def test_excel_all_columns_round_trip_without_shift(self):
        from openpyxl import Workbook, load_workbook

        row = {key: index + 0.25 for index, key in enumerate(user_flow.METRICS)}
        row.update(file="a.png", run_number=2, runs_per_file=2, actual_mode="text_only",
                   model_intervals_overlap=False, model_call_count=2,
                   layout_runtime_ms=None, text_detection_runtime_ms=0,
                   timing_missing_reasons={"layout_runtime_ms": "not_applicable"},
                   _model_calls=[{"call_id": name, "kind": "text_detection", "inference_ms": 0,
                                  "http_request_ms": 10, "model_load_ms": 0,
                                  "model_timings": {"text_detection": {"inference_ms": 0}}} for name in ("first", "second")])
        output = io.BytesIO()
        original_save = Workbook.save
        with patch.object(Path, "mkdir"), patch.object(Workbook, "save", autospec=True,
                side_effect=lambda workbook, path: original_save(workbook, output)):
            user_flow.save_excel(Path("results.xlsx"), [row])
        output.seek(0)
        workbook = load_workbook(output)
        try:
            summary = list(workbook["User Flow Timing"].values)
            self.assertEqual(tuple(summary[0]), user_flow.HEADERS)
            self.assertEqual(len(set(summary[0])), len(summary[0]))
            for key in user_flow.HEADERS:
                self.assertEqual(summary[1][summary[0].index(key)], user_flow.excel_value(row.get(key)), key)
            calls = list(workbook["Model Calls"].values)
            self.assertEqual(tuple(calls[0]), user_flow.MODEL_CALL_HEADERS)
            self.assertEqual(len(calls), 3)
            for exported, original in zip(calls[1:], row["_model_calls"]):
                expected = {**original, **{key: row[key] for key in ("file", "run_number", "actual_mode", "runs_per_file")}}
                for index, key in enumerate(calls[0]):
                    self.assertEqual(exported[index], user_flow.excel_value(expected.get(key)), key)
        finally:
            workbook.close()

    def arguments(self):
        return SimpleNamespace(repeats=2, frontend="http://localhost:3000", timeout_seconds=1200,
                               no_ocr=False, output=Path("results.xlsx"))

    def test_runs_each_file_twice_and_saves_after_every_run(self):
        cases = [(Path("a.png"), "standard"), (Path("b.pdf"), "standard")]
        rows = []
        saved_runs = []
        with patch.object(user_flow, "run_case", side_effect=lambda *args: {"status": "PASS", "ocr_status": "NOT_RUN"}) as run, patch.object(user_flow, "save_excel", side_effect=lambda path, results: saved_runs.append(len(results))), patch.object(user_flow, "reset_document") as reset, patch("builtins.print"):
            self.assertTrue(user_flow.run_cases(Mock(), self.arguments(), cases, rows, {"closed": False}))
        self.assertEqual([call.args[1].name for call in run.call_args_list], ["a.png", "a.png", "b.pdf", "b.pdf"])
        self.assertEqual([row["run_number"] for row in rows], [1, 2, 1, 2])
        self.assertTrue(all(row["runs_per_file"] == 2 for row in rows))
        self.assertEqual(saved_runs, [1, 2, 3, 4])
        self.assertEqual(reset.call_count, 4)

    def test_closed_page_becomes_failed_row_instead_of_traceback(self):
        page = Mock()
        page.evaluate.side_effect = RuntimeError("Browser context has been closed")
        row = user_flow.run_case(page, Path("a.png"), "standard", 1000)
        self.assertEqual(row["status"], "FAIL")
        self.assertIn("closed", row["error"])

    def test_browser_close_stops_remaining_runs_after_saving_result(self):
        state = {"closed": False}
        rows = []

        def close_browser(*args):
            state["closed"] = True
            return {"status": "FAIL", "ocr_status": "NOT_RUN"}

        with patch.object(user_flow, "run_case", side_effect=close_browser) as run, patch.object(user_flow, "save_excel") as save, patch("builtins.print"):
            self.assertFalse(user_flow.run_cases(Mock(), self.arguments(), [(Path("a.png"), None)], rows, state))
        run.assert_called_once()
        save.assert_called_once()
        self.assertEqual(len(rows), 1)

    def test_main_keeps_login_tab_open_until_all_cases_finish(self):
        context = Mock()
        login_page = Mock()
        context.pages = [login_page]
        driver = Mock()
        driver.chromium.launch_persistent_context.return_value = context
        manager = Mock()
        manager.__enter__ = Mock(return_value=driver)
        manager.__exit__ = Mock(return_value=False)
        api = SimpleNamespace(Error=RuntimeError, sync_playwright=lambda: manager)

        def verify_anchor(*args):
            login_page.close.assert_not_called()
            context.close.assert_not_called()
            return True

        with patch.dict("sys.modules", {"playwright": SimpleNamespace(), "playwright.sync_api": api}), patch("sys.argv", ["user_flow.py"]), patch.object(user_flow, "input_cases", return_value=[(Path("a.png"), None)]), patch.object(user_flow, "login_user") as login, patch.object(user_flow, "run_cases", side_effect=verify_anchor) as run, patch("builtins.print"):
            self.assertEqual(user_flow.main(), 0)
        login.assert_called_once_with(login_page, "http://localhost:3000")
        self.assertIs(run.call_args.args[0], login_page)
        login_page.close.assert_not_called()
        context.close.assert_called_once()

    def test_login_fills_requested_defaults(self):
        page = Mock()
        with patch.dict("os.environ", {}, clear=True):
            user_flow.login_user(page, "http://localhost:3000")
        page.goto.assert_called_once_with("http://localhost:3000/login?next=/", wait_until="domcontentloaded", timeout=30000)
        self.assertIn(unittest.mock.call('input[type="email"]'), page.locator.call_args_list)
        fills = [call.args[0] for call in page.locator.return_value.fill.call_args_list]
        self.assertEqual(fills, ["user@ocr.com", "user123"])

    def test_login_failure_does_not_expose_password(self):
        page = Mock()
        page.locator.return_value.fill.side_effect = RuntimeError("fill user123 failed")
        with self.assertRaises(RuntimeError) as failure:
            user_flow.login_user(page, "http://localhost:3000")
        self.assertNotIn("user123", str(failure.exception))

    def test_reset_waits_two_seconds_before_new_document(self):
        page = Mock()
        user_flow.reset_document(page)
        self.assertLess(page.mock_calls.index(unittest.mock.call.wait_for_timeout(2000)),
                        page.mock_calls.index(unittest.mock.call.get_by_role("button", name="เอกสารใหม่", exact=True)))
        page.locator.assert_called_once_with(user_flow.UPLOAD_INPUT)

    def test_found_and_no_match_stop_before_ocr_and_remove_listener(self):
        for matched in (True, False):
            with self.subTest(matched=matched):
                page = Mock()
                response = Mock()
                response.url = "http://localhost:18081/api/templates/detect-dev/jobs/test-job"
                response.json.return_value = {"data": {"status": "completed", "result": {
                    "matched": matched, "debug": {"verification_strategy": "standard"},
                }}}
                page.wait_for_function.side_effect = lambda *args, **kwargs: page.on.call_args.args[1](response)
                page.evaluate.side_effect = [None, [[{"metric": "end_to_end_ms", "valueMs": 123}]]]
                row = user_flow.run_case(page, Path("one.png"), "standard", 1000)
                self.assertEqual(row["status"], "PASS")
                self.assertEqual(row["matched"], matched)
                self.assertEqual(row["ocr_status"], "NOT_RUN")
                self.assertNotIn("ocr_wall_ms", row)
                self.assertEqual(page.get_by_role.call_count, 1)
                self.assertEqual(page.get_by_role.call_args.kwargs["name"], "ยืนยันเอกสารและค้นหา Template")
                page.remove_listener.assert_called_once_with("response", page.on.call_args.args[1])
                page.close.assert_not_called()

    @unittest.skipUnless(importlib.util.find_spec("openpyxl"), "openpyxl is needed for Excel round-trip test")
    def test_excel_sheets_keep_run_numbers(self):
        from openpyxl import Workbook, load_workbook

        output = io.BytesIO()
        original_save = Workbook.save
        rows = [{"file": "a.png", "run_number": run, "runs_per_file": 2, "actual_mode": "standard",
                 "_model_calls": [{"call_id": f"call-{run}"}]} for run in (1, 2)]
        with patch.object(Path, "mkdir"), patch.object(Workbook, "save", autospec=True, side_effect=lambda workbook, path: original_save(workbook, output)):
            user_flow.save_excel(Path("results.xlsx"), rows)
        output.seek(0)
        workbook = load_workbook(output)
        try:
            summary = list(workbook["User Flow Timing"].values)
            calls = list(workbook["Model Calls"].values)
            self.assertEqual([row[summary[0].index("run_number")] for row in summary[1:]], [1, 2])
            self.assertEqual([row[calls[0].index("run_number")] for row in calls[1:]], [1, 2])
        finally:
            workbook.close()

    def test_input_cases_group_by_mode(self):
        root = Path("input")
        files = [root / "standard" / "a.png", root / "strict" / "b.jpg", root / "ignore.txt"]
        with patch.object(Path, "exists", return_value=True), patch.object(Path, "rglob", return_value=files), patch.object(Path, "is_file", return_value=True):
            self.assertEqual([(path.name, mode) for path, mode in user_flow.input_cases(root, None)],
                             [("a.png", "standard"), ("b.jpg", "strict")])

    def test_mode_conflict_is_rejected(self):
        root = Path("input")
        with patch.object(Path, "exists", return_value=True), patch.object(Path, "rglob", return_value=[root / "strict" / "a.png"]), patch.object(Path, "is_file", return_value=True):
            with self.assertRaisesRegex(ValueError, "Mode conflict"):
                user_flow.input_cases(root, "standard")

    def test_completed_job_and_timing(self):
        result = {"matched": True, "best_candidate": {"verification_strategy": "standard"},
                  "debug": {"timing": {"backend_runtime_ms": 23, "total_detection_ms": 50}}}
        self.assertEqual(user_flow.detection_result({"data": {"status": "completed", "result": result}}), result)
        self.assertEqual(user_flow.actual_mode(result), "standard")
        values = user_flow.timing_values([[{"stage": "uploadJobMs", "elapsedMs": 5},
                                           {"metric": "end_to_end_ms", "valueMs": 80}]], result)
        self.assertEqual(values["uploadJobMs"], 5)
        self.assertEqual(values["end_to_end_ms"], 80)
        self.assertEqual(values["backend_runtime_ms"], 23)
        self.assertEqual(values["total_runtime_ms"], 50)
        self.assertIsNone(values["frontend_render_ms"])

    def test_excel_formula_is_not_evaluated(self):
        self.assertEqual(user_flow.excel_value("=HYPERLINK(\"x\")"), "'=HYPERLINK(\"x\")")

    def test_siglip_backend_measurement_replaces_empty_console_value(self):
        result = {"debug": {"timing": {"image_siglip_runtime_ms": 15.25}}}
        tables = [[{"metric": "image_siglip_runtime_ms", "valueMs": None}]]
        self.assertEqual(user_flow.timing_values(tables, result)["image_siglip_runtime_ms"], 15.25)
        legacy = {"debug": {"timing": {"siglip_inference_ms": 0.0}}}
        self.assertEqual(user_flow.timing_values(tables, legacy)["image_siglip_runtime_ms"], 0.0)
        self.assertIsNone(user_flow.timing_values(tables, {})["image_siglip_runtime_ms"])


if __name__ == "__main__":
    unittest.main()
